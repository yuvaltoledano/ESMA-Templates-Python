"""Taxonomy-driven check of the numeric-field set read as Float64.

`read_csv.read_and_clean` casts every present column whose taxonomy
FORMAT is {MONETARY}, {PERCENTAGE} or {INTEGER-n} to Float64 instead of
trusting Polars' type inference (config.NUMERIC_TAXONOMY_FORMATS,
config.INTEGER_TAXONOMY_FORMAT_PREFIX). The set is derived at run time
by `read_taxonomy.numeric_field_names`; this test re-derives it from the
bundled taxonomy independently, by reading the FORMAT cells directly,
and asserts the two agree exactly: every numeric-format field is in the
derived set and nothing else is.

Companion to test_date_columns_coverage.py, which plays the same role
for the {DATEFORMAT} fields.
"""

from __future__ import annotations

import openpyxl

from esma_milan.io_layer.clean_names import clean_name
from esma_milan.io_layer.read_taxonomy import (
    bundled_taxonomy_path,
    load_taxonomy,
    load_taxonomy_formats,
    numeric_field_names,
)

TAXONOMY_PATH = bundled_taxonomy_path()

# Sheet 1 column indices (0-indexed) in the bundled taxonomy:
#   0 TEMPLATE CATEGORY, 1 SECTION, 2 FIELD CODE, 3 FIELD NAME,
#   4 CONTENT TO REPORT, 5 ND1-ND4, 6 ND5, 7 FORMAT
_FIELD_CODE_COL = 2
_FIELD_NAME_COL = 3
_FORMAT_COL = 7

# The bundled taxonomy carries 19 {MONETARY}, 12 {PERCENTAGE} and
# 5 {INTEGER-9999} fields.
_EXPECTED_NUMERIC_FIELD_COUNT = 36


def _taxonomy_rows() -> list[tuple[str, str, str]]:
    """Return [(field_code, cleaned_field_name, format)] for every data row."""
    wb = openpyxl.load_workbook(TAXONOMY_PATH, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        ws.reset_dimensions()
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    out: list[tuple[str, str, str]] = []
    for row in rows[1:]:  # skip header
        if len(row) <= _FORMAT_COL:
            continue
        code = row[_FIELD_CODE_COL]
        name = row[_FIELD_NAME_COL]
        fmt = row[_FORMAT_COL]
        if not (isinstance(code, str) and isinstance(name, str) and isinstance(fmt, str)):
            continue
        out.append((code.strip(), clean_name(name.strip()), fmt.strip()))
    return out


def _expected_numeric_names() -> set[str]:
    return {
        name
        for _, name, fmt in _taxonomy_rows()
        if fmt in ("{MONETARY}", "{PERCENTAGE}") or fmt.startswith("{INTEGER-")
    }


def test_taxonomy_numeric_fields_extracted() -> None:
    """Sanity: the bundled taxonomy yields the expected number of numeric
    fields, with at least one of each format. If this fails, the
    extraction shape changed and the coverage check below is meaningless."""
    rows = _taxonomy_rows()
    formats = {fmt for _, _, fmt in rows}
    assert "{MONETARY}" in formats
    assert "{PERCENTAGE}" in formats
    assert any(f.startswith("{INTEGER-") for f in formats)
    assert len(_expected_numeric_names()) == _EXPECTED_NUMERIC_FIELD_COUNT


def test_derived_numeric_set_matches_taxonomy_exactly() -> None:
    """Every {MONETARY} / {PERCENTAGE} / {INTEGER-n} field's cleaned name is
    in the derived set, and nothing else is."""
    taxonomy = load_taxonomy(TAXONOMY_PATH)
    formats = load_taxonomy_formats(TAXONOMY_PATH)
    derived = numeric_field_names(taxonomy, formats)
    expected = _expected_numeric_names()

    missing = sorted(expected - derived)
    extra = sorted(derived - expected)
    assert not missing, f"numeric taxonomy fields missing from derived set: {missing}"
    assert not extra, f"non-numeric fields present in derived set: {extra}"
    assert len(derived) == _EXPECTED_NUMERIC_FIELD_COUNT


def test_structural_balance_and_valuation_fields_are_numeric() -> None:
    """The four columns the pipeline uses in arithmetic before Stage 9
    must be in the set; a taxonomy edit that dropped them would re-open
    the sample-inference crash on exactly the fields that matter most."""
    derived = numeric_field_names(load_taxonomy(TAXONOMY_PATH), load_taxonomy_formats(TAXONOMY_PATH))
    for col in (
        "current_principal_balance",
        "original_principal_balance",
        "current_valuation_amount",
        "original_valuation_amount",
    ):
        assert col in derived, col


def test_date_list_and_numeric_set_are_disjoint() -> None:
    """A field cannot be both parsed as a date and cast to Float64."""
    from esma_milan.config import LOAN_DATE_COLUMNS, PROPERTY_DATE_COLUMNS

    derived = numeric_field_names(load_taxonomy(TAXONOMY_PATH), load_taxonomy_formats(TAXONOMY_PATH))
    overlap = derived & (set(LOAN_DATE_COLUMNS) | set(PROPERTY_DATE_COLUMNS))
    assert not overlap, overlap
