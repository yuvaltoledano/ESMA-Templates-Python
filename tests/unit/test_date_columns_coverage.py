"""Taxonomy-driven coverage check for LOAN_DATE_COLUMNS / PROPERTY_DATE_COLUMNS.

The Excel-serial encoding in the workbook writer detects date columns by
Polars dtype (`pl.Date`). Stage 1 promotes a String/Int column to
`pl.Date` only when its cleaned name appears in LOAN_DATE_COLUMNS or
PROPERTY_DATE_COLUMNS. So a taxonomy {DATEFORMAT} field missing from
either tuple silently slips through as ISO-string or raw-int in the
final workbook — exactly the Domi-2025-1 parity failure for
pool_addition_date, principal_grace_period_end_date, and
prepayment_date.

This test reads the bundled synthetic taxonomy and asserts every
{DATEFORMAT}-typed RREL/RREC field's cleaned name is present in the
corresponding date-columns tuple. If a future taxonomy update introduces
a new date field without updating the parse list, this test fails and
points the maintainer at the gap.

Companion to `test_write_workbook_emits_excel_serials_for_pl_date_columns`
in test_write_workbook.py: the writer test catches "column got dropped
between parse and write"; this test catches "new taxonomy date field
added without updating parse list".
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from esma_milan.config import (
    ALWAYS_DROPPED_COLUMNS,
    LOAN_DATE_COLUMNS,
    PROPERTY_DATE_COLUMNS,
)
from esma_milan.io_layer.clean_names import clean_name

REPO_ROOT = Path(__file__).resolve().parents[2]
TAXONOMY_PATH = REPO_ROOT / "tests" / "fixtures" / "synthetic" / "taxonomy.xlsx"

# Sheet 1 column indices (0-indexed) in the bundled taxonomy:
#   0 TEMPLATE CATEGORY, 1 SECTION, 2 FIELD CODE, 3 FIELD NAME,
#   4 CONTENT TO REPORT, 5 ND1-ND4, 6 ND5, 7 FORMAT
_FIELD_CODE_COL = 2
_FIELD_NAME_COL = 3
_FORMAT_COL = 7

_DATE_FORMAT_TOKEN = "DATEFORMAT"


def _taxonomy_date_fields() -> list[tuple[str, str]]:
    """Return [(field_code, cleaned_field_name)] for every taxonomy row
    whose FORMAT cell mentions {DATEFORMAT}."""
    wb = openpyxl.load_workbook(TAXONOMY_PATH, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        ws.reset_dimensions()
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    out: list[tuple[str, str]] = []
    for row in rows[1:]:  # skip header
        if len(row) <= _FORMAT_COL:
            continue
        code = row[_FIELD_CODE_COL]
        name = row[_FIELD_NAME_COL]
        fmt = row[_FORMAT_COL]
        if not (isinstance(code, str) and isinstance(name, str) and isinstance(fmt, str)):
            continue
        if _DATE_FORMAT_TOKEN not in fmt:
            continue
        out.append((code.strip(), clean_name(name.strip())))
    return out


def test_taxonomy_date_fields_extracted() -> None:
    """Sanity: the bundled taxonomy yields at least one RREL date field
    and at least one RREC date field. If this assertion ever fails, the
    extraction shape changed and the coverage check below is meaningless."""
    fields = _taxonomy_date_fields()
    rrel = [c for c, _ in fields if c.startswith("RREL")]
    rrec = [c for c, _ in fields if c.startswith("RREC")]
    assert rrel, "no RREL {DATEFORMAT} fields found in taxonomy"
    assert rrec, "no RREC {DATEFORMAT} fields found in taxonomy"


def test_every_taxonomy_loan_date_field_in_loan_date_columns() -> None:
    """Every RREL {DATEFORMAT} field's cleaned name must appear in
    LOAN_DATE_COLUMNS (or in ALWAYS_DROPPED_COLUMNS, which excuses
    `data_cut_off_date` — RREL6 is metadata dropped before Stage 1's
    date-parse loop)."""
    loan_dates = set(LOAN_DATE_COLUMNS)
    dropped = set(ALWAYS_DROPPED_COLUMNS)
    missing: list[tuple[str, str]] = []
    for code, cleaned in _taxonomy_date_fields():
        if not code.startswith("RREL"):
            continue
        if cleaned in dropped:
            continue
        if cleaned not in loan_dates:
            missing.append((code, cleaned))
    assert not missing, (
        "Taxonomy {DATEFORMAT} RREL fields missing from LOAN_DATE_COLUMNS:\n"
        + "\n".join(f"  {c} -> {n!r}" for c, n in missing)
        + "\nAdd them to LOAN_DATE_COLUMNS in src/esma_milan/config.py "
        "(or to ALWAYS_DROPPED_COLUMNS if they are metadata)."
    )


def test_every_taxonomy_property_date_field_in_property_date_columns() -> None:
    """Every RREC {DATEFORMAT} field's cleaned name must appear in
    PROPERTY_DATE_COLUMNS (or in ALWAYS_DROPPED_COLUMNS)."""
    property_dates = set(PROPERTY_DATE_COLUMNS)
    dropped = set(ALWAYS_DROPPED_COLUMNS)
    missing: list[tuple[str, str]] = []
    for code, cleaned in _taxonomy_date_fields():
        if not code.startswith("RREC"):
            continue
        if cleaned in dropped:
            continue
        if cleaned not in property_dates:
            missing.append((code, cleaned))
    assert not missing, (
        "Taxonomy {DATEFORMAT} RREC fields missing from PROPERTY_DATE_COLUMNS:\n"
        + "\n".join(f"  {c} -> {n!r}" for c, n in missing)
        + "\nAdd them to PROPERTY_DATE_COLUMNS in src/esma_milan/config.py "
        "(or to ALWAYS_DROPPED_COLUMNS if they are metadata)."
    )
